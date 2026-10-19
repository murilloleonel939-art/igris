"""Generación de documentos Word (.docx) bajo pedido del usuario.

Por qué existe este módulo
--------------------------
Antes, la única forma de "escribir en Word" era `type_in_office`, que pega texto
con Ctrl+V en la ventana que ya estuviera activa. Eso no crea un documento: no
abre Word, no guarda nada y no devuelve ningún nombre de archivo. Cuando el
usuario pedía una presentación, el modelo rellenaba ese vacío afirmando que el
documento estaba listo, y era falso.

Aquí el archivo se construye de verdad en el servidor, se guarda con un nombre
concreto y se devuelve un enlace de descarga. Si algo falla, se informa del
fallo; nunca se dice que existe algo que no se ha escrito.

El documento queda asociado a su dueño (user_id). El enlace de descarga
comprueba ese dueño, así que un identificador ajeno no sirve para leer el
documento de otra persona.
"""
from __future__ import annotations

import io
import logging
import os
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

logger = logging.getLogger(__name__)

# Carpeta donde se guardan los .docx. En Docker debe ser un volumen persistente
# si se quieren conservar entre despliegues; si se pierde, el usuario puede
# volver a pedir el documento.
DOCUMENTS_DIR = Path(
    os.environ.get(
        "IGRIS_DOCUMENTS_DIR",
        str(Path(__file__).resolve().parent.parent / "documents"),
    )
)

# Tope de texto por campo. Un documento enorme no cabe en el contexto del modelo
# y tampoco es lo que el usuario ha pedido.
MAX_BODY_CHARS = 6000
MAX_SECTIONS = 40
MAX_TITLE_CHARS = 160


def _available() -> bool:
    """True si python-docx está instalado."""
    try:
        import docx  # noqa: F401
        return True
    except Exception:
        return False


def _slug(text: str) -> str:
    """Nombre de archivo seguro: sin acentos raros, sin rutas, sin espacios."""
    normalized = unicodedata.normalize("NFKD", text or "")
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    cleaned = re.sub(r"[^A-Za-z0-9]+", "-", ascii_text).strip("-").lower()
    return (cleaned or "documento")[:60]


def _clean(value: Any, limit: int = MAX_BODY_CHARS) -> str:
    return str(value or "").strip()[:limit]


def _parse_sections(raw: Any) -> list[dict[str, Any]]:
    """Normaliza las secciones que envía el modelo.

    Acepta tanto dicts ({heading, body}) como texto suelto, porque el modelo a
    veces simplifica. Lo que no encaje se descarta en vez de guardarse mal.
    """
    if not isinstance(raw, list):
        return []
    sections: list[dict[str, Any]] = []
    for item in raw[:MAX_SECTIONS]:
        if isinstance(item, str):
            text = _clean(item)
            if text:
                sections.append({"heading": "", "body": [text], "bullets": False})
            continue
        if not isinstance(item, dict):
            continue
        heading = _clean(item.get("heading") or item.get("title"), MAX_TITLE_CHARS)
        body = item.get("body")
        if isinstance(body, str):
            lines = [line.strip() for line in body.splitlines() if line.strip()]
        elif isinstance(body, list):
            lines = [_clean(line) for line in body]
            lines = [line for line in lines if line]
        else:
            lines = []
        if not heading and not lines:
            continue
        sections.append({
            "heading": heading,
            "body": lines,
            "bullets": bool(item.get("bullets")),
        })
    return sections


def build_docx(
    title: str,
    subtitle: str = "",
    author: str = "",
    sections: Optional[list] = None,
    cover: bool = True,
) -> bytes:
    """Construye el .docx en memoria y devuelve los bytes.

    Lanza RuntimeError si falta python-docx; el llamador traduce eso a un mensaje
    claro para el usuario en lugar de generar un archivo vacío.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
        from docx.shared import Pt
    except Exception as exc:  # pragma: no cover - depende del despliegue
        raise RuntimeError(
            "Falta el paquete python-docx en el servidor. Instálalo para poder "
            "crear documentos Word."
        ) from exc

    doc = Document()
    parsed = _parse_sections(sections)

    # Márgenes y tipografía sobrios: un documento legible sin plantillas.
    for section in doc.sections:
        section.top_margin = Pt(72)
        section.bottom_margin = Pt(72)
        section.left_margin = Pt(72)
        section.right_margin = Pt(72)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    if cover and title:
        # Portada: título centrado, subtítulo y fecha. Luego salto de página.
        for _ in range(6):
            doc.add_paragraph()
        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(title)
        run.bold = True
        run.font.size = Pt(28)
        if subtitle:
            sub = doc.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub_run = sub.add_run(subtitle)
            sub_run.italic = True
            sub_run.font.size = Pt(14)
        for _ in range(4):
            doc.add_paragraph()
        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run(
            (f"{author}\n" if author else "")
            + datetime.now().strftime("%d/%m/%Y")
        )
        footer_run.font.size = Pt(10)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    elif title:
        doc.add_heading(title, level=0)

    for section in parsed:
        if section["heading"]:
            doc.add_heading(section["heading"], level=1)
        for line in section["body"]:
            if section["bullets"]:
                doc.add_paragraph(line, style="List Bullet")
            else:
                doc.add_paragraph(line)

    if not parsed and not title:
        doc.add_paragraph("Documento vacío.")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def save_document(
    user_id: int,
    title: str,
    content: bytes,
    filename: str,
) -> Path:
    """Escribe el archivo en disco, un directorio por usuario."""
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    folder = DOCUMENTS_DIR / f"u{user_id}"
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / filename
    target.write_bytes(content)
    return target


def make_filename(title: str) -> str:
    """Nombre visible y único: titulo-AAAAMMDD-HHMMSS.docx"""
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return f"{_slug(title)}-{stamp}.docx"


def section_count(sections: Optional[list]) -> int:
    return len(_parse_sections(sections))
